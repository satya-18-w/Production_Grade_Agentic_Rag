"""
Lightweight true_data parser for building/verifying the golden dataset.

Deliberately independent of app/ingestion/ — uses python-docx and
BeautifulSoup directly rather than the production loaders or `unstructured`,
so this stays fast and has no dependency on the ingestion pipeline's own
correctness. The golden dataset's `relevant_contexts` were sourced by
reading these same files directly; this script makes that reproducible
rather than a one-off manual read.
"""

import os

TRUE_DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "DATA", "true_data")

SOURCE_FILES = {
    "parallel_work_queue": "parallel_work_queue.txt",
    "pods_autoscale": "pods_autoscale.html",
    "job_management": "job_management.html",
    "cronjobs": "cronjobs.docx",
    "monitor_job": "monitor_job.docx",
}


def parse_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def parse_html(path: str) -> str:
    from bs4 import BeautifulSoup

    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text(separator="\n")


def parse_docx(path: str) -> str:
    import docx

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def parse_true_data_file(domain: str) -> str:
    """Parse one true_data source by its golden-dataset domain name."""
    filename = SOURCE_FILES[domain]
    path = os.path.join(TRUE_DATA_DIR, filename)
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "txt":
        return parse_txt(path)
    if ext in ("html", "htm"):
        return parse_html(path)
    if ext == "docx":
        return parse_docx(path)
    raise ValueError(f"Unsupported extension for golden dataset parsing: {ext}")


def parse_all() -> dict[str, str]:
    return {domain: parse_true_data_file(domain) for domain in SOURCE_FILES}


if __name__ == "__main__":
    # Dump each source's parsed text — useful for verifying a
    # relevant_contexts entry in golden_dataset.json actually appears
    # verbatim (or near-verbatim) in the real source document.
    for domain, text in parse_all().items():
        print("=" * 80)
        print(domain, f"({len(text)} chars)")
        print("=" * 80)
        print(text[:500])
        print("...\n")
