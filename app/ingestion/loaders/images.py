import io
import ipaddress
import os
import socket
from typing import NamedTuple, List, Optional, Tuple
from urllib.parse import urlparse

import logfire

from app.config import settings


class ExtractedImage(NamedTuple):
    data: bytes
    ext: str
    mime_type: str
    location: str  # human-readable locator, e.g. "page 3", "slide 5", "image 2"


# Skip tiny images — icons, bullets, logos, avatars, decorative dividers.
# These add captioning cost/noise without carrying real technical content.
# Conservative threshold; real diagrams/screenshots are almost always well
# above this (a confirmed 64x64 author-avatar in pods_autoscale.html is
# exactly the kind of thing this is meant to filter out).
MIN_IMAGE_BYTES = 3000

# External fetch safety limits — Kubernetes docs sourced from web articles
# (Medium, vendor docs) commonly keep images as external links rather than
# embedding them (confirmed: 20 of 20 images in cronjobs.docx, all 4 images
# across the true_data HTML files). Fetching them is opt-out, not opt-in,
# but stays bounded: public hosts only, size-capped, timeout-bounded.
MAX_FETCH_BYTES = 15 * 1024 * 1024  # 15 MB — generous for a diagram, blocks abuse
FETCH_TIMEOUT = 10  # seconds
_FETCH_USER_AGENT = "Mozilla/5.0 (compatible; EnterpriseRAGIngestion/1.0)"

_EXT_TO_MIME = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "gif": "image/gif",
    "bmp": "image/bmp",
    "tiff": "image/tiff",
    "webp": "image/webp",
    "svg": "image/svg+xml",
}
_MIME_TO_EXT = {v: k for k, v in _EXT_TO_MIME.items() if k != "jpg"}


# ── PDF ──────────────────────────────────────────────────────────────────────

def extract_images_from_pdf(file_path: str) -> List[ExtractedImage]:
    """
    Extract embedded raster images from a PDF using pypdf's built-in image
    support (page.images) — no new dependency; pypdf handles the filter
    decoding internally and returns ready-to-write bytes per image.
    """
    from pypdf import PdfReader

    images: List[ExtractedImage] = []
    with logfire.span("Image Extraction (PDF)", filename=file_path):
        try:
            reader = PdfReader(file_path)
            for page_num, page in enumerate(reader.pages, start=1):
                for img in page.images:
                    if len(img.data) < MIN_IMAGE_BYTES:
                        continue
                    ext = img.name.rsplit(".", 1)[-1].lower() if "." in img.name else "png"
                    mime = _EXT_TO_MIME.get(ext, "image/png")
                    images.append(ExtractedImage(data=img.data, ext=ext, mime_type=mime, location=f"page {page_num}"))
            logfire.info(f"Extracted {len(images)} image(s) from {file_path}.")
        except Exception as e:
            # Best-effort — a broken image stream shouldn't fail text extraction.
            logfire.warning(f"Image extraction failed for {file_path}: {e}")
    return images


# ── PPTX ─────────────────────────────────────────────────────────────────────

def extract_images_from_pptx(file_path: str) -> List[ExtractedImage]:
    """
    Extract embedded images from a PPTX using python-pptx's shape.image API.
    Note: slides built from native vector shapes (freeform/auto-shape, common
    for hand-drawn architecture diagrams) have no PICTURE shapes at all and
    will legitimately yield zero images — that's not a bug, there's simply
    nothing to extract without rendering the slide itself.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    images: List[ExtractedImage] = []
    with logfire.span("Image Extraction (PPTX)", filename=file_path):
        try:
            prs = Presentation(file_path)
            for slide_num, slide in enumerate(prs.slides, start=1):
                for shape in slide.shapes:
                    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                        continue
                    image = shape.image
                    if len(image.blob) < MIN_IMAGE_BYTES:
                        continue
                    images.append(ExtractedImage(
                        data=image.blob,
                        ext=image.ext,
                        mime_type=image.content_type,
                        location=f"slide {slide_num}",
                    ))
            logfire.info(f"Extracted {len(images)} image(s) from {file_path}.")
        except Exception as e:
            logfire.warning(f"Image extraction failed for {file_path}: {e}")
    return images


# ── DOCX ─────────────────────────────────────────────────────────────────────

def extract_images_from_docx(file_path: str) -> List[ExtractedImage]:
    """
    Extract images referenced in a DOCX — both embedded and externally-linked.
    Docs sourced from web articles (Medium, vendor docs) tend to keep images
    as external links rather than embedding them, so external fetch is a
    first-class path here, not a fallback.
    """
    from docx import Document as DocxDocument

    images: List[ExtractedImage] = []
    with logfire.span("Image Extraction (DOCX)", filename=file_path):
        try:
            doc = DocxDocument(file_path)
            img_rels = [r for r in doc.part.rels.values() if "image" in r.reltype]

            for idx, rel in enumerate(img_rels, start=1):
                if rel.is_external:
                    fetched = _fetch_external_image(rel.target_ref)
                else:
                    try:
                        fetched = (rel.target_part.blob, rel.target_part.content_type)
                    except Exception:
                        fetched = None

                if fetched is None:
                    continue
                data, mime = fetched
                if len(data) < MIN_IMAGE_BYTES:
                    continue
                ext = _MIME_TO_EXT.get(mime, mime.split("/")[-1] if "/" in mime else "png")
                images.append(ExtractedImage(data=data, ext=ext, mime_type=mime, location=f"image {idx}"))

            logfire.info(f"Extracted {len(images)} image(s) from {file_path}.")
        except Exception as e:
            logfire.warning(f"Image extraction failed for {file_path}: {e}")
    return images


# ── HTML ─────────────────────────────────────────────────────────────────────

def extract_images_from_html(file_path: str) -> List[ExtractedImage]:
    """
    Extract images referenced by <img> tags in a local HTML file — data URIs
    decoded directly, external URLs fetched with safety limits, relative
    paths resolved against the HTML file's own directory.
    """
    from bs4 import BeautifulSoup

    images: List[ExtractedImage] = []
    base_dir = os.path.dirname(os.path.abspath(file_path))

    with logfire.span("Image Extraction (HTML)", filename=file_path):
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                soup = BeautifulSoup(f.read(), "html.parser")

            for idx, img in enumerate(soup.find_all("img"), start=1):
                src = (img.get("src") or "").strip()
                if not src:
                    continue

                if src.startswith("data:"):
                    fetched = _decode_data_uri(src)
                elif src.startswith("http://") or src.startswith("https://"):
                    fetched = _fetch_external_image(src)
                else:
                    fetched = _read_local_image(os.path.join(base_dir, src), base_dir)

                if fetched is None:
                    continue
                data, mime = fetched
                if len(data) < MIN_IMAGE_BYTES:
                    continue
                ext = _MIME_TO_EXT.get(mime, mime.split("/")[-1] if "/" in mime else "png")
                images.append(ExtractedImage(data=data, ext=ext, mime_type=mime, location=f"image {idx}"))

            logfire.info(f"Extracted {len(images)} image(s) from {file_path}.")
        except Exception as e:
            logfire.warning(f"Image extraction failed for {file_path}: {e}")
    return images


# ── Shared helpers: external fetch (SSRF-guarded), data URI, local path ────────

def _is_safe_public_host(hostname: str) -> bool:
    """
    Basic SSRF guard for the external-image fetch path: reject hosts that
    resolve to private / loopback / link-local / reserved / multicast IPs.
    Image URLs come from document content (docx/html source files chosen for
    ingestion), not live user input — but the same hygiene applies regardless
    of source.
    """
    try:
        addr_info = socket.getaddrinfo(hostname, None)
    except socket.gaierror:
        return False
    for info in addr_info:
        try:
            ip = ipaddress.ip_address(info[4][0])
        except ValueError:
            continue
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            return False
    return True


def _fetch_external_image(url: str) -> Optional[Tuple[bytes, str]]:
    """
    Fetch an external image URL with safety limits: http/https + public-host
    only, size-capped streaming download, request timeout, content-type
    validated as image/*. Returns (data, mime_type) or None on any
    failure/rejection — best-effort, never raises.
    """
    if not settings.ALLOW_EXTERNAL_IMAGE_FETCH:
        return None

    try:
        import requests

        parsed = urlparse(url)
        if parsed.scheme not in ("http", "https") or not parsed.hostname:
            return None
        if not _is_safe_public_host(parsed.hostname):
            logfire.warning(f"Refusing to fetch image from non-public host: {url}")
            return None

        resp = requests.get(
            url,
            timeout=FETCH_TIMEOUT,
            stream=True,
            headers={"User-Agent": _FETCH_USER_AGENT},
        )
        resp.raise_for_status()

        content_type = resp.headers.get("Content-Type", "").split(";")[0].strip().lower()
        if not content_type.startswith("image/"):
            return None

        chunks = []
        total = 0
        for chunk in resp.iter_content(chunk_size=65536):
            total += len(chunk)
            if total > MAX_FETCH_BYTES:
                logfire.warning(f"Image fetch exceeded {MAX_FETCH_BYTES} byte cap, aborting: {url}")
                return None
            chunks.append(chunk)

        return b"".join(chunks), content_type
    except Exception as e:
        logfire.warning(f"External image fetch failed ({url}): {e}")
        return None


def _decode_data_uri(data_uri: str) -> Optional[Tuple[bytes, str]]:
    """Decode a base64 `data:` URI image into (data, mime_type)."""
    import base64

    try:
        header, b64data = data_uri.split(",", 1)
        mime = (header.split(";")[0].replace("data:", "").strip() or "image/png")
        return base64.b64decode(b64data), mime
    except Exception:
        return None


def _read_local_image(path: str, base_dir: str) -> Optional[Tuple[bytes, str]]:
    """
    Read a locally-referenced image (relative <img src> path) from disk.
    Resolves symlinks/`..` and confirms the result stays inside base_dir
    before opening it — a relative src in a source HTML file shouldn't be
    able to walk outside the directory it was found in.
    """
    try:
        real_base = os.path.realpath(base_dir)
        real_path = os.path.realpath(path)
        if real_path != real_base and not real_path.startswith(real_base + os.sep):
            return None
        if not os.path.isfile(real_path):
            return None
        with open(real_path, "rb") as f:
            data = f.read()
        ext = real_path.rsplit(".", 1)[-1].lower() if "." in real_path else "png"
        return data, _EXT_TO_MIME.get(ext, "image/png")
    except Exception:
        return None


# ── PPTX vector-shape diagram rendering (Track E) ──────────────────────────
#
# extract_images_from_pptx() above only finds embedded PICTURE shapes.
# Slides built entirely from native vector shapes (FREEFORM/AUTO_SHAPE —
# common for hand-drawn architecture diagrams, confirmed on all 66 slides
# of DATA/true_data/architecture.pptx) have nothing for it to find, since
# there is no embedded raster image to extract — the diagram only exists
# as vector geometry. This section renders those specific slides to PNG
# via headless LibreOffice instead. Genuinely optional: everything else in
# this file works with zero system dependencies; this is the one exception,
# and it degrades to "no rendered slides" rather than failing ingestion if
# LibreOffice isn't installed.

RENDER_TIMEOUT = 60  # seconds — a single-file headless conversion, generous but bounded

_SOFFICE_CANDIDATES = [
    "soffice",
    "libreoffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
]


def _find_soffice() -> Optional[str]:
    """
    Locate a LibreOffice/soffice binary — checks PATH first via shutil.which,
    then a few common absolute install locations across platforms, since
    LibreOffice frequently isn't added to PATH by its own installer
    (particularly on Windows and macOS). Returns None if not found anywhere
    — the caller treats this as "feature unavailable," not an error.
    """
    import shutil

    for candidate in _SOFFICE_CANDIDATES:
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
        if os.path.isfile(candidate):
            return candidate
    return None


def needs_slide_render(slide) -> bool:
    """
    A slide is a rendering candidate only if it has FREEFORM/AUTO_SHAPE
    content (likely a diagram) AND no PICTURE shape was already found on
    it — avoids rendering purely-text slides (redundant with the existing
    text extraction) and slides that already have a real embedded image.
    Pure python-pptx shape inspection — no LibreOffice needed for this
    check, so it's cheap to run on every slide before deciding whether the
    expensive conversion step is worth invoking at all.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    has_diagram_shapes = False
    has_picture = False
    for shape in slide.shapes:
        if shape.shape_type in (MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.AUTO_SHAPE):
            has_diagram_shapes = True
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            has_picture = True
    return has_diagram_shapes and not has_picture


def render_pptx_diagram_slides(file_path: str) -> List[ExtractedImage]:
    """
    Render vector-shape diagram slides (per needs_slide_render) to PNG via
    headless LibreOffice, for slides where extract_images_from_pptx() finds
    nothing because there's no embedded picture to extract.

    Best-effort at every step, consistent with every other extractor in
    this file: no candidate slides, no LibreOffice binary, a conversion
    timeout, or any subprocess/rendering failure all degrade to an empty
    list rather than raising — a missing system dependency should never
    fail the rest of ingestion.
    """
    import subprocess
    import tempfile
    from pptx import Presentation

    images: List[ExtractedImage] = []

    with logfire.span("Slide Rendering (PPTX diagrams)", filename=file_path):
        try:
            prs = Presentation(file_path)
            candidate_slide_numbers = [
                i for i, slide in enumerate(prs.slides, start=1)
                if needs_slide_render(slide)
            ]
        except Exception as e:
            logfire.warning(f"Could not inspect slides for rendering candidates in {file_path}: {e}")
            return images

        if not candidate_slide_numbers:
            return images

        soffice = _find_soffice()
        if not soffice:
            logfire.warning(
                f"{len(candidate_slide_numbers)} diagram slide(s) found in {file_path} but "
                "LibreOffice is not installed — skipping slide rendering. "
                "See DOCS/13_MULTIMODAL_RAG.md for the system dependency."
            )
            return images

        try:
            from pdf2image import convert_from_path
        except ImportError:
            logfire.warning(
                "pdf2image is not installed — skipping slide rendering "
                "(LibreOffice was found, but the PDF page renderer is missing)."
            )
            return images

        with tempfile.TemporaryDirectory() as tmp_dir:
            try:
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp_dir, file_path],
                    timeout=RENDER_TIMEOUT,
                    capture_output=True,
                    check=True,
                )
            except Exception as e:
                logfire.warning(f"LibreOffice conversion failed for {file_path}: {e}")
                return images

            pdf_candidates = [f for f in os.listdir(tmp_dir) if f.lower().endswith(".pdf")]
            if not pdf_candidates:
                logfire.warning(f"LibreOffice produced no PDF output for {file_path}.")
                return images
            pdf_path = os.path.join(tmp_dir, pdf_candidates[0])

            for slide_num in candidate_slide_numbers:
                try:
                    pages = convert_from_path(pdf_path, first_page=slide_num, last_page=slide_num, dpi=150)
                    if not pages:
                        continue
                    buf = io.BytesIO()
                    pages[0].save(buf, format="PNG")
                    data = buf.getvalue()
                    if len(data) < MIN_IMAGE_BYTES:
                        continue
                    images.append(ExtractedImage(
                        data=data, ext="png", mime_type="image/png",
                        location=f"slide {slide_num} (rendered)",
                    ))
                except Exception as e:
                    logfire.warning(f"Could not render slide {slide_num} of {file_path}: {e}")

        logfire.info(f"Rendered {len(images)} diagram slide(s) from {file_path}.")

    return images
